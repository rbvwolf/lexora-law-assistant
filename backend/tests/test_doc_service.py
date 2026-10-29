"""doc_service tests: real .docx rendering, template integrity and KVKK leakage checks.

Imports doc_service directly (no FastAPI app), so these run without network, Groq or
sentence-transformers.
"""

import asyncio
import json
import os
import re
import threading
import uuid
import zipfile
from datetime import datetime, timedelta, timezone
from pathlib import Path

import pytest
from docx import Document
from docxtpl import DocxTemplate
from pydantic import ValidationError
from pypdf import PdfReader

from app.core.config import Settings
from app.services import doc_service as doc_service_module
from app.services.doc_layouts import FOOTER_TEXT, get_layout
from app.services.doc_service import KIRA_ARTIS_TEMPLATE_ID, DocumentAssetMissingError, doc_service

FIXTURES_DIR = Path(__file__).parent / "fixtures"

UUID_DOCX_RE = re.compile(r"^[0-9a-f]{8}-[0-9a-f]{4}-4[0-9a-f]{3}-[89ab][0-9a-f]{3}-[0-9a-f]{12}\.docx$")


def _kira_artis_case_params():
    return {
        "kira_baslangic_tarihi": "2022-09-01",
        "mevcut_kira_tutari": 12000,
        "talep_edilen_kira_tutari": 24000,
        "tufe_oran": 62.51,
        "noter_ili": "Ankara",
    }


def _kira_artis_pii():
    return {
        "kiraci_ad_soyad": "Tolga Yılmaz",
        "kiraci_tc": "19482910382",
        "kiraci_adres": "Ayrancı Mah. Güvenlik Cad. 24/5 Çankaya / ANKARA",
        "kiraya_veren_ad_soyad": "Hikmet Çetin",
        "kiraya_veren_tc": "38291048214",
        "kiraya_veren_adres": "Kızılay Mah. Atatürk Bulvarı 10/3 Çankaya / ANKARA",
        "tasinmaz_adresi": "Bahçelievler Mah. 7. Cadde 12/4 Çankaya / ANKARA",
    }


def _document_text(path) -> str:
    doc = Document(str(path))
    parts = [p.text for p in doc.paragraphs]
    for section in doc.sections:
        parts.extend(p.text for p in section.footer.paragraphs)
    return "\n".join(parts)


@pytest.fixture
def temp_storage(tmp_path, monkeypatch):
    monkeypatch.setattr(doc_service_module, "STORAGE_TEMP_DIR", tmp_path)
    return tmp_path


async def _generate(case_params, pii, template_id=KIRA_ARTIS_TEMPLATE_ID):
    template = doc_service.get_template(template_id)
    errors = doc_service.validate_fields(template, case_params, pii)
    assert not errors.has_errors, errors.to_message()
    return await doc_service.generate_documents(template_id, case_params, pii)


# ---------------------------------------------------------------------------
# Template integrity
# ---------------------------------------------------------------------------


def test_template_file_exists_for_kira_artis():
    assert doc_service.template_path(KIRA_ARTIS_TEMPLATE_ID).is_file()


def test_template_tags_are_not_split_across_runs():
    """docxtpl only replaces tags that sit inside one <w:t>; a split tag leaves '{{' behind."""
    with zipfile.ZipFile(doc_service.template_path(KIRA_ARTIS_TEMPLATE_ID)) as zf:
        xml = zf.read("word/document.xml").decode("utf-8")
    texts = re.findall(r"<w:t(?:\s[^>]*)?>([^<]*)</w:t>", xml)
    assert any("{{" in t for t in texts)
    for text in texts:
        assert text.count("{{") == text.count("}}"), f"Split tag in run: {text!r}"
        assert text.count("{%") == text.count("%}"), f"Split block tag in run: {text!r}"


def test_template_variables_match_render_context():
    tpl = DocxTemplate(str(doc_service.template_path(KIRA_ARTIS_TEMPLATE_ID)))
    declared = tpl.get_undeclared_template_variables()
    context = doc_service.build_render_context(KIRA_ARTIS_TEMPLATE_ID, _kira_artis_case_params(), _kira_artis_pii())
    assert declared == set(context)


def test_template_listed_with_expected_fields():
    template = doc_service.get_template(KIRA_ARTIS_TEMPLATE_ID)
    case_names = {f.field_name for f in template.case_param_fields}
    pii_names = {f.field_name for f in template.pii_fields}
    assert case_names == {"kira_baslangic_tarihi", "mevcut_kira_tutari", "talep_edilen_kira_tutari", "tufe_oran", "noter_ili"}
    assert {"kiraci_ad_soyad", "kiraci_adres", "kiraya_veren_adres", "tasinmaz_adresi"} <= pii_names
    # Addresses are PII and must never sit in the LLM-facing case_params
    assert not any("adres" in name for name in case_names)


def test_existing_templates_unchanged():
    assert doc_service.get_template("kira_ihtarname_temerrut") is not None
    assert doc_service.get_template("tuketici_ayipli_mal_ihtarname") is not None


# ---------------------------------------------------------------------------
# Rendering
# ---------------------------------------------------------------------------


@pytest.mark.asyncio
async def test_generate_real_docx_replaces_all_fields(temp_storage):
    pii = _kira_artis_pii()
    file_id, docx_url, _, _ = await _generate(_kira_artis_case_params(), pii)

    docx_path = temp_storage / f"{file_id}.docx"
    assert docx_url.endswith(f"{file_id}.docx")
    text = _document_text(docx_path)

    assert "{{" not in text and "}}" not in text
    for value in pii.values():
        assert value in text
    assert "Ankara Nöbetçi Noterliğine" in text
    assert "01.09.2022" in text
    assert "12.000,00 TL" in text
    assert "24.000,00 TL" in text
    assert "%100,00" in text
    assert "avukatlık hizmeti vermez" in text
    assert not list(temp_storage.glob("*.tmp"))


@pytest.mark.asyncio
async def test_legal_ceiling_computed_from_user_rate(temp_storage):
    file_id, *_ = await _generate(_kira_artis_case_params(), _kira_artis_pii())
    text = _document_text(temp_storage / f"{file_id}.docx")
    # 12.000 x (1 + 62,51 / 100) = 19.501,20
    assert "%62,51" in text
    assert "19.501,20 TL" in text


@pytest.mark.asyncio
async def test_rate_is_not_hardcoded(temp_storage):
    case_params = _kira_artis_case_params()
    case_params["tufe_oran"] = "40"
    file_id, *_ = await _generate(case_params, _kira_artis_pii())
    text = _document_text(temp_storage / f"{file_id}.docx")
    assert "%40,00" in text
    assert "16.800,00 TL" in text
    assert "62,51" not in text


@pytest.mark.asyncio
async def test_turkish_and_xml_special_characters_survive(temp_storage):
    pii = _kira_artis_pii()
    pii["kiraya_veren_ad_soyad"] = "Çağlar & Oğulları <Şirketi> İşçi Ğüşıö"
    file_id, *_ = await _generate(_kira_artis_case_params(), pii)
    text = _document_text(temp_storage / f"{file_id}.docx")
    assert "Çağlar & Oğulları <Şirketi> İşçi Ğüşıö" in text
    assert "KEŞİDECİ" in text and "İHTARNAME" in text


@pytest.mark.asyncio
async def test_optional_tc_omitted_cleanly(temp_storage):
    pii = _kira_artis_pii()
    del pii["kiraci_tc"]
    del pii["kiraya_veren_tc"]
    file_id, *_ = await _generate(_kira_artis_case_params(), pii)
    text = _document_text(temp_storage / f"{file_id}.docx")
    assert "T.C. Kimlik No" not in text
    assert "None" not in text


@pytest.mark.asyncio
async def test_dropped_case_param_not_in_document(temp_storage):
    case_params = _kira_artis_case_params()
    case_params["llm_ek_alan"] = "gizli_deger_123"
    file_id, *_ = await _generate(case_params, _kira_artis_pii())
    assert "gizli_deger_123" not in _document_text(temp_storage / f"{file_id}.docx")


# ---------------------------------------------------------------------------
# KVKK: file names, mock files, logs
# ---------------------------------------------------------------------------


@pytest.mark.asyncio
async def test_file_names_are_random_and_pii_free(temp_storage):
    pii = _kira_artis_pii()
    file_id, docx_url, pdf_url, _ = await _generate(_kira_artis_case_params(), pii)

    uuid.UUID(file_id, version=4)
    names = [p.name for p in temp_storage.iterdir()]
    assert sorted(names) == sorted([f"{file_id}.docx", f"{file_id}.pdf"])
    assert UUID_DOCX_RE.match(f"{file_id}.docx")
    for value in pii.values():
        for token in value.split():
            if len(token) > 3:
                assert token not in docx_url and token not in pdf_url
                assert all(token not in name for name in names)


@pytest.mark.asyncio
async def test_mock_files_do_not_contain_pii(temp_storage):
    case_params = {
        "kira_tutari": 18000,
        "gecikilen_aylar": "Eylül-Ekim 2026",
        "odeme_suresi_gun": 30,
        "tasinmaz_adresi": "Kadıköy / İstanbul",
    }
    pii = {
        "kiraya_veren_ad_soyad": "Mehmet Yılmaz",
        "kiraya_veren_tc": "11111111111",
        "kiraci_ad_soyad": "Ali Demir",
        "iban": "TR120006100511123456789012",
    }
    file_id, *_ = await _generate(case_params, pii, template_id="kira_ihtarname_temerrut")
    for suffix in (".docx", ".pdf"):
        content = (temp_storage / f"{file_id}{suffix}").read_text(encoding="utf-8")
        assert "kira_ihtarname_temerrut" in content
        for value in list(pii.values()) + ["Kadıköy", "18000"]:
            assert value not in content


@pytest.mark.asyncio
async def test_generation_logs_no_pii(temp_storage, caplog):
    pii = _kira_artis_pii()
    with caplog.at_level("DEBUG"):
        await _generate(_kira_artis_case_params(), pii)
    for value in pii.values():
        assert value not in caplog.text


@pytest.mark.asyncio
async def test_missing_template_file_falls_back_to_mock_in_development(temp_storage, tmp_path, monkeypatch):
    empty_dir = tmp_path / "no_templates"
    empty_dir.mkdir()
    monkeypatch.setattr(doc_service_module.settings, "TEMPLATES_DIR", str(empty_dir))
    monkeypatch.setattr(doc_service_module.settings, "ENVIRONMENT", "development")
    pii = _kira_artis_pii()
    file_id, *_ = await _generate(_kira_artis_case_params(), pii)
    content = (temp_storage / f"{file_id}.docx").read_text(encoding="utf-8")
    assert KIRA_ARTIS_TEMPLATE_ID in content
    assert pii["kiraci_ad_soyad"] not in content


# ---------------------------------------------------------------------------
# Validation
# ---------------------------------------------------------------------------


def _validate(case_params, pii):
    template = doc_service.get_template(KIRA_ARTIS_TEMPLATE_ID)
    return doc_service.validate_fields(template, case_params, pii)


def test_valid_payload_passes():
    assert not _validate(_kira_artis_case_params(), _kira_artis_pii()).has_errors


@pytest.mark.parametrize("field", ["tufe_oran", "mevcut_kira_tutari", "noter_ili", "kira_baslangic_tarihi"])
def test_missing_required_case_param(field):
    case_params = _kira_artis_case_params()
    del case_params[field]
    errors = _validate(case_params, _kira_artis_pii())
    assert errors.missing == [field]


@pytest.mark.parametrize("field", ["kiraci_adres", "kiraya_veren_adres", "tasinmaz_adresi", "kiraci_ad_soyad"])
def test_missing_required_pii(field):
    pii = _kira_artis_pii()
    del pii[field]
    errors = _validate(_kira_artis_case_params(), pii)
    assert errors.missing == [field]


@pytest.mark.parametrize("talep", [12000, 15000, 19501.20, "19501.2"])
def test_request_within_legal_ceiling_rejected(talep):
    case_params = _kira_artis_case_params()
    case_params["talep_edilen_kira_tutari"] = talep
    errors = _validate(case_params, _kira_artis_pii())
    assert errors.invalid == ["talep_edilen_kira_tutari"]
    message = errors.to_message()
    assert "talep_edilen_kira_tutari" in message
    assert str(talep) not in message
    assert "19501" not in message and "19.501" not in message


def test_request_one_kurus_over_ceiling_accepted():
    case_params = _kira_artis_case_params()
    case_params["talep_edilen_kira_tutari"] = 19501.21
    assert not _validate(case_params, _kira_artis_pii()).has_errors


@pytest.mark.parametrize("oran", [-1, 200.01, 6251, "nan", "inf"])
def test_tufe_rate_out_of_range_rejected(oran):
    case_params = _kira_artis_case_params()
    case_params["tufe_oran"] = oran
    errors = _validate(case_params, _kira_artis_pii())
    assert errors.invalid == ["tufe_oran"]
    assert str(oran) not in errors.to_message()


@pytest.mark.parametrize("mevcut", [0, -500])
def test_non_positive_current_rent_rejected(mevcut):
    case_params = _kira_artis_case_params()
    case_params["mevcut_kira_tutari"] = mevcut
    errors = _validate(case_params, _kira_artis_pii())
    assert errors.invalid == ["mevcut_kira_tutari"]


def test_cross_field_rule_skipped_when_field_already_invalid():
    case_params = _kira_artis_case_params()
    case_params["talep_edilen_kira_tutari"] = "yirmi dört bin"
    errors = _validate(case_params, _kira_artis_pii())
    assert errors.invalid == ["talep_edilen_kira_tutari"]


def test_invalid_tc_rejected_without_value_in_message():
    pii = _kira_artis_pii()
    pii["kiraci_tc"] = "1948291038"
    errors = _validate(_kira_artis_case_params(), pii)
    assert errors.invalid == ["kiraci_tc"]
    assert "1948291038" not in errors.to_message()


@pytest.mark.parametrize("environment", ["staging", "production"])
@pytest.mark.asyncio
async def test_missing_template_file_raises_outside_development(temp_storage, tmp_path, monkeypatch, environment):
    empty_dir = tmp_path / "no_templates"
    empty_dir.mkdir()
    monkeypatch.setattr(doc_service_module.settings, "TEMPLATES_DIR", str(empty_dir))
    monkeypatch.setattr(doc_service_module.settings, "ENVIRONMENT", environment)
    pii = _kira_artis_pii()
    with pytest.raises(DocumentAssetMissingError) as exc_info:
        await doc_service.generate_documents(KIRA_ARTIS_TEMPLATE_ID, _kira_artis_case_params(), pii)
    message = str(exc_info.value)
    assert "docx_template" in message
    assert str(empty_dir) not in message and "no_templates" not in message
    for value in pii.values():
        assert value not in message
    assert not [p for p in temp_storage.iterdir() if p.is_file()]


@pytest.mark.parametrize("environment", ["staging", "production"])
@pytest.mark.asyncio
async def test_missing_fonts_raise_outside_development(temp_storage, tmp_path, monkeypatch, environment):
    monkeypatch.setattr(doc_service_module.settings, "FONTS_DIR", str(tmp_path / "no_fonts"))
    monkeypatch.setattr(doc_service_module.settings, "ENVIRONMENT", environment)
    with pytest.raises(DocumentAssetMissingError) as exc_info:
        await doc_service.generate_documents(KIRA_ARTIS_TEMPLATE_ID, _kira_artis_case_params(), _kira_artis_pii())
    assert "pdf_fonts" in str(exc_info.value)
    assert "no_fonts" not in str(exc_info.value)
    assert not [p for p in temp_storage.iterdir() if p.is_file()]


@pytest.mark.asyncio
async def test_missing_fonts_fall_back_to_mock_pdf_in_development(temp_storage, tmp_path, monkeypatch):
    monkeypatch.setattr(doc_service_module.settings, "FONTS_DIR", str(tmp_path / "no_fonts"))
    monkeypatch.setattr(doc_service_module.settings, "ENVIRONMENT", "development")
    pii = _kira_artis_pii()
    file_id, *_ = await _generate(_kira_artis_case_params(), pii)
    content = (temp_storage / f"{file_id}.pdf").read_text(encoding="utf-8")
    assert content.startswith("%PDF-1.4 Mock PDF")
    assert pii["kiraci_ad_soyad"] not in content
    # The .docx is still real
    assert pii["kiraci_ad_soyad"] in _document_text(temp_storage / f"{file_id}.docx")


@pytest.mark.asyncio
async def test_render_failure_leaves_no_files(temp_storage, monkeypatch):
    def broken_pdf(*args, **kwargs):
        raise ValueError("render failed")

    monkeypatch.setattr(doc_service, "_render_pdf", broken_pdf)
    with pytest.raises(ValueError):
        await doc_service.generate_documents(KIRA_ARTIS_TEMPLATE_ID, _kira_artis_case_params(), _kira_artis_pii())
    assert list(temp_storage.iterdir()) == []


# ---------------------------------------------------------------------------
# Single layout source: .docx unchanged by the refactor, PDF built from the same layout
# ---------------------------------------------------------------------------


def _paragraph_snapshot(p):
    return {
        "text": p.text,
        "alignment": str(p.alignment),
        "first_line_indent": p.paragraph_format.first_line_indent,
        "runs": [[r.text, bool(r.bold)] for r in p.runs],
    }


def test_rendered_docx_matches_pre_refactor_snapshot(tmp_path):
    """Same input as Desktop/ornek_ihtarname.docx (rendered before doc_layouts existed)."""
    expected = json.loads((FIXTURES_DIR / "ornek_ihtarname_docx.json").read_text(encoding="utf-8"))
    pii = {
        "kiraci_ad_soyad": "Deniz Örnekoğlu",
        "kiraci_tc": "12345678901",
        "kiraci_adres": "Örnek Mah. Deneme Sok. No: 5/3 Çankaya / ANKARA",
        "kiraya_veren_ad_soyad": "Şebnem Kurgusal",
        "kiraya_veren_adres": "Hayali Mah. Taslak Cad. No: 18/1 Yenimahalle / ANKARA",
        "tasinmaz_adresi": "Örnek Mah. Deneme Sok. No: 5/3 Çankaya / ANKARA",
    }
    context = doc_service.build_render_context(KIRA_ARTIS_TEMPLATE_ID, _kira_artis_case_params(), pii)
    context["belge_tarihi"] = "29.10.2026"
    out = tmp_path / "snapshot.docx"
    doc_service._render_docx(doc_service.template_path(KIRA_ARTIS_TEMPLATE_ID), context, out)

    doc = Document(str(out))
    assert [_paragraph_snapshot(p) for p in doc.paragraphs] == expected["paragraphs"]
    assert [_paragraph_snapshot(p) for p in doc.sections[0].footer.paragraphs] == expected["footer"]


def test_committed_docx_template_matches_layout():
    """Fails if doc_layouts changed but tools/build_doc_templates was not re-run."""
    layout = get_layout(KIRA_ARTIS_TEMPLATE_ID)
    doc = Document(str(doc_service.template_path(KIRA_ARTIS_TEMPLATE_ID)))
    actual = [[[r.text, bool(r.bold)] for r in p.runs] for p in doc.paragraphs]
    expected = [[[text, bold] for text, bold in spec.segments] for spec in layout.paragraphs]
    assert actual == expected
    assert doc.sections[0].footer.paragraphs[0].text == layout.footer == FOOTER_TEXT


# ---------------------------------------------------------------------------
# PDF
# ---------------------------------------------------------------------------


def _pdf_text(path) -> str:
    reader = PdfReader(str(path))
    raw = "\n".join(page.extract_text() for page in reader.pages)
    # Justified lines get stretched spaces and line breaks; compare on single spaces
    return re.sub(r"\s+", " ", raw)


@pytest.mark.asyncio
async def test_pdf_contains_all_values_and_turkish_characters(temp_storage):
    pii = _kira_artis_pii()
    pii["kiraci_adres"] = "Öğretmenler Sk. Işık Apt. Üçüncü Kat 4/7 Çankaya / İSTANBUL ğüşıöçĞÜŞİÖÇ"
    file_id, *_ = await _generate(_kira_artis_case_params(), pii)
    pdf_path = temp_storage / f"{file_id}.pdf"
    assert pdf_path.read_bytes().startswith(b"%PDF-")
    text = _pdf_text(pdf_path)

    for value in pii.values():
        assert value in text
    for expected in (
        "Ankara Nöbetçi Noterliğine",
        "İHTARNAME",
        "KEŞİDECİ (Kiracı):",
        "MUHATAP (Kiraya Veren):",
        "01.09.2022",
        "12.000,00 TL",
        "24.000,00 TL",
        "%100,00",
        "%62,51",
        "19.501,20 TL",
        "SONUÇ:",
        "İmza",
    ):
        assert expected in text
    for char in "ğüşıöçĞÜŞİÖÇ":
        assert char in text
    assert "{{" not in text


@pytest.mark.asyncio
async def test_pdf_keeps_special_characters_literal(temp_storage):
    pii = _kira_artis_pii()
    pii["kiraya_veren_ad_soyad"] = "Çağlar & Oğulları <Şirketi> **kalın değil** İşçi Ğüşıö"
    file_id, *_ = await _generate(_kira_artis_case_params(), pii)
    assert pii["kiraya_veren_ad_soyad"] in _pdf_text(temp_storage / f"{file_id}.pdf")


@pytest.mark.asyncio
async def test_pdf_footer_disclaimer(temp_storage):
    file_id, *_ = await _generate(_kira_artis_case_params(), _kira_artis_pii())
    assert FOOTER_TEXT in _pdf_text(temp_storage / f"{file_id}.pdf")


@pytest.mark.asyncio
async def test_pdf_metadata_has_no_pii(temp_storage):
    pii = _kira_artis_pii()
    file_id, *_ = await _generate(_kira_artis_case_params(), pii)
    metadata = PdfReader(str(temp_storage / f"{file_id}.pdf")).metadata
    assert metadata.title == "Fahiş Kira Artışına İtiraz İhtarnamesi"
    assert metadata.author == "lexora"
    flat = " ".join(str(v) for v in metadata.values())
    for value in pii.values():
        for token in value.split():
            if len(token) > 3:
                assert token not in flat


@pytest.mark.asyncio
async def test_docx_and_pdf_show_same_computed_values(temp_storage):
    case_params = _kira_artis_case_params()
    case_params["tufe_oran"] = "33.33"
    case_params["talep_edilen_kira_tutari"] = 17500.5
    file_id, *_ = await _generate(case_params, _kira_artis_pii())
    docx_text = re.sub(r"\s+", " ", _document_text(temp_storage / f"{file_id}.docx"))
    pdf_text = _pdf_text(temp_storage / f"{file_id}.pdf")
    # 12.000 x 1,3333 = 15.999,60 ; (17.500,50 - 12.000) / 12.000 = %45,84
    for expected in ("15.999,60 TL", "%33,33", "17.500,50 TL", "%45,84"):
        assert expected in docx_text
        assert expected in pdf_text


@pytest.mark.asyncio
async def test_mocked_templates_still_get_mock_pdf(temp_storage):
    case_params = {
        "urun_adi": "Buzdolabı",
        "fatura_tarihi": "2026-09-15",
        "ayipli_durum_ozeti": "Soğutmuyor",
        "secimlik_hak": "Değişim",
    }
    pii = {"tuketici_ad_soyad": "Ayşe Kaya", "tuketici_tc": "11111111111", "satici_unvan": "Örnek AŞ"}
    file_id, *_ = await _generate(case_params, pii, template_id="tuketici_ayipli_mal_ihtarname")
    content = (temp_storage / f"{file_id}.pdf").read_text(encoding="utf-8")
    assert content == "%PDF-1.4 Mock PDF Document for template tuketici_ayipli_mal_ihtarname"


# ---------------------------------------------------------------------------
# Expiry (fake clock, never real waiting)
# ---------------------------------------------------------------------------


class FakeClock:
    def __init__(self, start: datetime):
        self.now = start

    def __call__(self) -> datetime:
        return self.now

    def advance(self, **kwargs) -> None:
        self.now += timedelta(**kwargs)


@pytest.fixture
def clock(monkeypatch):
    fake = FakeClock(datetime(2026, 10, 29, 12, 0, tzinfo=timezone.utc))
    monkeypatch.setattr(doc_service_module, "_now", fake)
    return fake


@pytest.mark.asyncio
async def test_expires_at_equals_mtime_plus_ttl(temp_storage, clock, monkeypatch):
    monkeypatch.setattr(doc_service_module.settings, "DOCUMENT_TTL_MINUTES", 45)
    file_id, _, _, expires_at = await _generate(_kira_artis_case_params(), _kira_artis_pii())
    assert expires_at == clock.now + timedelta(minutes=45)
    for suffix in (".docx", ".pdf"):
        mtime = (temp_storage / f"{file_id}{suffix}").stat().st_mtime
        assert mtime + 45 * 60 == pytest.approx(expires_at.timestamp())


@pytest.mark.asyncio
async def test_download_allowed_until_expiry_then_deleted(temp_storage, clock):
    file_id, *_ = await _generate(_kira_artis_case_params(), _kira_artis_pii())
    docx_name = f"{file_id}.docx"

    clock.advance(minutes=120, seconds=-1)
    assert doc_service.resolve_download_path(docx_name) == temp_storage / docx_name

    clock.advance(seconds=1)
    assert doc_service.resolve_download_path(docx_name) is None
    assert not (temp_storage / docx_name).exists()


def test_tmp_and_other_suffixes_are_never_served(temp_storage):
    for name in ("abc.docx.tmp", "abc.pdf.tmp", "abc.txt", "abc"):
        (temp_storage / name).write_text("x", encoding="utf-8")
        assert doc_service.resolve_download_path(name) is None


def _touch(path: Path, mtime: datetime) -> None:
    path.write_text("x", encoding="utf-8")
    os.utime(path, (mtime.timestamp(), mtime.timestamp()))


def test_sweep_removes_only_expired_files_including_tmp(temp_storage, clock):
    ttl = timedelta(minutes=120)
    old = clock.now - ttl
    fresh = clock.now - ttl + timedelta(seconds=1)
    _touch(temp_storage / "old.docx", old)
    _touch(temp_storage / "old.pdf", old)
    _touch(temp_storage / "old.docx.tmp", old)
    _touch(temp_storage / "fresh.docx", fresh)
    _touch(temp_storage / "fresh.pdf.tmp", fresh)
    subdir = temp_storage / "subdir"
    subdir.mkdir()
    os.utime(subdir, (old.timestamp(), old.timestamp()))

    assert doc_service.sweep_expired() == 3
    assert sorted(p.name for p in temp_storage.iterdir()) == ["fresh.docx", "fresh.pdf.tmp", "subdir"]


def test_sweep_handles_files_left_from_before_restart(temp_storage, clock):
    """Expiry comes from mtime, so nothing in-memory is needed after a restart."""
    _touch(temp_storage / "leftover.pdf", clock.now - timedelta(days=3))
    assert doc_service.sweep_expired() == 1


def test_sweep_with_missing_storage_dir(tmp_path, monkeypatch):
    monkeypatch.setattr(doc_service_module, "STORAGE_TEMP_DIR", tmp_path / "missing")
    assert doc_service.sweep_expired() == 0


@pytest.mark.asyncio
async def test_sweep_async_runs_in_worker_thread(monkeypatch):
    seen = {}

    def fake_sweep():
        seen["thread"] = threading.current_thread()
        return 0

    monkeypatch.setattr(doc_service, "sweep_expired", fake_sweep)
    await doc_service.sweep_expired_async()
    assert seen["thread"] is not threading.main_thread()


@pytest.mark.asyncio
async def test_renderers_run_in_worker_threads(temp_storage, monkeypatch):
    seen = []
    real_docx, real_pdf = doc_service._render_docx, doc_service._render_pdf

    def spy_docx(*args):
        seen.append(threading.current_thread())
        return real_docx(*args)

    def spy_pdf(*args):
        seen.append(threading.current_thread())
        return real_pdf(*args)

    monkeypatch.setattr(doc_service, "_render_docx", spy_docx)
    monkeypatch.setattr(doc_service, "_render_pdf", spy_pdf)
    await _generate(_kira_artis_case_params(), _kira_artis_pii())
    assert len(seen) == 2
    assert all(t is not threading.main_thread() for t in seen)


@pytest.mark.asyncio
async def test_sweeper_loop_survives_errors_and_cancels(monkeypatch):
    calls = []

    async def flaky_sweep():
        calls.append(1)
        raise OSError("disk")

    monkeypatch.setattr(doc_service, "sweep_expired_async", flaky_sweep)
    monkeypatch.setattr(doc_service_module, "DOCUMENT_SWEEP_INTERVAL", timedelta(seconds=0))
    task = asyncio.create_task(doc_service.run_sweeper())
    for _ in range(5):
        await asyncio.sleep(0)
    task.cancel()
    with pytest.raises(asyncio.CancelledError):
        await task
    assert len(calls) >= 2


def test_sweep_interval_is_ten_minutes():
    assert doc_service_module.DOCUMENT_SWEEP_INTERVAL == timedelta(minutes=10)


def test_ttl_configurable_via_env(monkeypatch):
    monkeypatch.setenv("DOCUMENT_TTL_MINUTES", "5")
    assert Settings().DOCUMENT_TTL_MINUTES == 5
    monkeypatch.setenv("DOCUMENT_TTL_MINUTES", "0")
    with pytest.raises(ValidationError):
        Settings()


def test_default_ttl_is_two_hours():
    assert Settings.model_fields["DOCUMENT_TTL_MINUTES"].default == 120
